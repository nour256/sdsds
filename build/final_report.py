"""Print the final slide list, per-slide notes word counts and the claim->timestamp map."""
import json, re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation('IEN_301_Project_1_v2.pptx')
rep = json.load(open('build_report.json'))

def spoken(txt):
    return len(re.findall(r"[A-Za-z0-9’'\-–]+", txt.split('\n\nClick:')[0]))

print('SLIDE LIST')
total = 0
for n, s in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    lab = texts[0] if texts else ''
    head = ''
    for t in texts[1:]:
        if len(t) > 6 and not t[:2].isdigit():
            head = t.replace('\n', ' / ')
            break
    hidden = ' (hidden)' if s._element.get('show') == '0' else ''
    w = spoken(rep['notes'][str(n)])
    total += w if n <= 16 else 0
    print(f'{n:2d}{hidden:9s} {lab[:44]:44s} | {head[:70]:70s} | notes {w:3d} w')
print(f'\nSpoken notes total (slides 1-16): {total} words = {total/130:.1f} min at 130 wpm (hidden slide 17 excluded)')
print('\nCLAIM -> RECORDING MAP')
seen = set()
for sl, claim, iv, ts, verb in rep['claims']:
    k = (claim, iv, ts)
    if k in seen:
        continue
    seen.add(k)
    print(f'  slide {sl:2d} | {claim[:48]:48s} | Interview {iv} @ {ts} | "{verb}"')
print('\nNUMBERS:', json.dumps(rep['numbers']))
