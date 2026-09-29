"""Helpers for rebuilding IEN301 deck in the exact style of the original slides 2-10.

Theme values were read from the original file (orig.pptx, Google Slides export):
  slide size 9144000 x 5143500 EMU (10in x 5.625in)
  font            Quattrocento Sans (embedded in the package)
  light bg        F5F7F9   dark bg  05080D
  ink (titles)    111820   white    FFFFFF
  accent blue     006FE8   sky blue (dark slides) 80D6FF
  grey body       52606D   mist body (dark slides) B5C1CC
  table rule      D7DFE5
  section label   (0.62, 0.39, 8.75, 0.23) in, 11pt bold, accent
  footer caveat   (0.62, 5.24, 8.36, 0.20) in, 10pt, grey/mist
  slide number    (9.14, 5.24, 0.39, 0.21) in, 10pt, grey/mist
"""
import copy
from lxml import etree
from pptx.util import Emu, Inches, Pt
from pptx.oxml.ns import qn

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
NS = {'a': A, 'p': P, 'r': R, 'mc': MC, 'p14': P14}

FONT = 'Quattrocento Sans'
DARK_BG, LIGHT_BG = '05080D', 'F5F7F9'
INK, WHITE = '111820', 'FFFFFF'
BLUE, SKY = '006FE8', '80D6FF'
GREY, MIST = '52606D', 'B5C1CC'
RULE = 'D7DFE5'
CARD = 'FFFFFF'
TINT = 'E6F1FD'   # very light blue used only as a subtle card/icon-circle tint on light slides

SLIDE_W, SLIDE_H = 9144000, 5143500


def emu(v):
    """inches -> EMU int"""
    return int(round(v * 914400))


def E(tag, attrib=None, *children, text=None):
    ns, local = tag.split(':')
    el = etree.Element('{%s}%s' % (NS[ns], local), nsmap=None)
    if attrib:
        for k, v in attrib.items():
            if ':' in k:
                kns, kl = k.split(':')
                el.set('{%s}%s' % (NS[kns], kl), str(v))
            else:
                el.set(k, str(v))
    for c in children:
        if c is not None:
            el.append(c)
    if text is not None:
        el.text = text
    return el


# ----------------------------------------------------------------------------- text
def rpr(sz, color, bold=False, italic=False, font=FONT):
    attrs = {'lang': 'en', 'sz': str(int(sz * 100)), 'b': '1' if bold else '0'}
    if italic:
        attrs['i'] = '1'
    return E('a:rPr', attrs,
             E('a:solidFill', None, E('a:srgbClr', {'val': color})),
             E('a:latin', {'typeface': font}), E('a:ea', {'typeface': font}),
             E('a:cs', {'typeface': font}), E('a:sym', {'typeface': font}))


def run(text, sz, color, bold=False, italic=False):
    t = E('a:t', None, text=text)
    if text != text.strip():
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return E('a:r', None, rpr(sz, color, bold, italic), t)


def para(runs, algn='l', spc_bef=0, spc_aft=0, ln_spc=None, bullet=None, mar_l=0, indent=0):
    """runs: list of a:r elements (or (text, sz, color, bold) tuples)."""
    ppr = E('a:pPr', {'indent': str(indent), 'lvl': '0', 'marL': str(mar_l), 'marR': '0', 'rtl': '0', 'algn': algn})
    if ln_spc:
        ppr.append(E('a:lnSpc', None, E('a:spcPct', {'val': str(int(ln_spc * 100000))})))
    ppr.append(E('a:spcBef', None, E('a:spcPts', {'val': str(int(spc_bef * 100))})))
    ppr.append(E('a:spcAft', None, E('a:spcPts', {'val': str(int(spc_aft * 100))})))
    if bullet:
        ppr.append(E('a:buClr', None, E('a:srgbClr', {'val': bullet[1]})))
        ppr.append(E('a:buSzPts', {'val': str(int(bullet[2] * 100))}))
        ppr.append(E('a:buFont', {'typeface': FONT}))
        ppr.append(E('a:buChar', {'char': bullet[0]}))
    else:
        ppr.append(E('a:buNone'))
    p = E('a:p', None, ppr)
    for r in runs:
        if isinstance(r, tuple):
            r = run(*r)
        p.append(r)
    return p


def next_id(slide):
    ids = [int(e.get('id')) for e in slide._element.iter('{%s}cNvPr' % P)]
    return (max(ids) + 1) if ids else 2


def _sp_tree(slide):
    return slide.shapes._spTree


def textbox(slide, x, y, w, h, paras, anchor='t', name=None, wrap='square', autofit=False, insets=(0, 0, 0, 0), fill=None, line=None, geom='rect', rot=None):
    """Add a text shape. paras: list of a:p elements (from para()). Returns shape id."""
    sid = next_id(slide)
    name = name or f'TextBox {sid}'
    xfrm = E('a:xfrm', {'rot': str(rot)} if rot else None, E('a:off', {'x': emu(x), 'y': emu(y)}), E('a:ext', {'cx': emu(w), 'cy': emu(h)}))
    sppr = E('p:spPr', None, xfrm, E('a:prstGeom', {'prst': geom}, E('a:avLst')))
    sppr.append(E('a:solidFill', None, E('a:srgbClr', {'val': fill})) if fill else E('a:noFill'))
    if line:
        sppr.append(E('a:ln', {'w': str(line[1])}, E('a:solidFill', None, E('a:srgbClr', {'val': line[0]}))))
    else:
        sppr.append(E('a:ln', None, E('a:noFill')))
    l, t, r, b = insets
    bodypr = E('a:bodyPr', {'anchorCtr': '0', 'anchor': anchor, 'bIns': emu(b), 'lIns': emu(l), 'spcFirstLastPara': '1', 'rIns': emu(r), 'wrap': wrap, 'tIns': emu(t)},
               E('a:normAutofit') if autofit else E('a:noAutofit'))
    txbody = E('p:txBody', None, bodypr, E('a:lstStyle'))
    for p in paras:
        txbody.append(p)
    sp = E('p:sp', None,
           E('p:nvSpPr', None, E('p:cNvPr', {'id': sid, 'name': name}), E('p:cNvSpPr', {'txBox': '1'}), E('p:nvPr')),
           sppr, txbody)
    _sp_tree(slide).append(sp)
    return sid


def shape(slide, geom, x, y, w, h, fill=None, line=None, alpha=None, rot=None, name=None, adj=None, dash=None):
    """Plain auto shape (no text). line=(color, width_emu). alpha = 0..100 percent opacity of fill."""
    sid = next_id(slide)
    name = name or f'Shape {sid}'
    xfrm = E('a:xfrm', {'rot': str(rot)} if rot else None, E('a:off', {'x': emu(x), 'y': emu(y)}), E('a:ext', {'cx': emu(w), 'cy': emu(h)}))
    av = E('a:avLst')
    if adj:
        for k, v in adj.items():
            av.append(E('a:gd', {'name': k, 'fmla': f'val {v}'}))
    sppr = E('p:spPr', None, xfrm, E('a:prstGeom', {'prst': geom}, av))
    if fill:
        clr = E('a:srgbClr', {'val': fill})
        if alpha is not None:
            clr.append(E('a:alpha', {'val': str(int(alpha * 1000))}))
        sppr.append(E('a:solidFill', None, clr))
    else:
        sppr.append(E('a:noFill'))
    if line:
        ln = E('a:ln', {'w': str(line[1]), 'cap': 'rnd'}, E('a:solidFill', None, E('a:srgbClr', {'val': line[0]})))
        if dash:
            ln.append(E('a:prstDash', {'val': dash}))
        ln.append(E('a:round'))
        sppr.append(ln)
    else:
        sppr.append(E('a:ln', None, E('a:noFill')))
    sp = E('p:sp', None,
           E('p:nvSpPr', None, E('p:cNvPr', {'id': sid, 'name': name}), E('p:cNvSpPr'), E('p:nvPr')),
           sppr,
           E('p:txBody', None, E('a:bodyPr', {'anchor': 'ctr', 'lIns': '0', 'rIns': '0', 'tIns': '0', 'bIns': '0'}, E('a:noAutofit')), E('a:lstStyle'), para([])))
    _sp_tree(slide).append(sp)
    return sid


def line(slide, x1, y1, x2, y2, color, width_emu=7150, name=None):
    """Straight connector line."""
    sid = next_id(slide)
    name = name or f'Line {sid}'
    x, y = min(x1, x2), min(y1, y2)
    w, h = abs(x2 - x1), abs(y2 - y1)
    flip = {}
    if x2 < x1:
        flip['flipH'] = '1'
    if y2 < y1:
        flip['flipV'] = '1'
    xfrm = E('a:xfrm', flip, E('a:off', {'x': emu(x), 'y': emu(y)}), E('a:ext', {'cx': emu(w), 'cy': emu(h)}))
    cxn = E('p:cxnSp', None,
            E('p:nvCxnSpPr', None, E('p:cNvPr', {'id': sid, 'name': name}), E('p:cNvCxnSpPr'), E('p:nvPr')),
            E('p:spPr', None, xfrm, E('a:prstGeom', {'prst': 'straightConnector1'}, E('a:avLst')),
              E('a:ln', {'w': str(width_emu), 'cap': 'rnd'}, E('a:solidFill', None, E('a:srgbClr', {'val': color})), E('a:round'))))
    _sp_tree(slide).append(cxn)
    return sid


def picture(slide, path, x, y, w, h, crop=None, to_back=True, name=None):
    """crop=(left, top, right, bottom) fractions 0..1"""
    pic = slide.shapes.add_picture(path, emu(x), emu(y), emu(w), emu(h))
    if name:
        pic.name = name
    if crop:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    if to_back:
        tree = _sp_tree(slide)
        tree.remove(pic._element)
        tree.insert(2, pic._element)
    return pic.shape_id


def set_background(slide, color):
    cSld = slide._element.find('{%s}cSld' % P)
    old = cSld.find('{%s}bg' % P)
    if old is not None:
        cSld.remove(old)
    bg = E('p:bg', None, E('p:bgPr', None, E('a:solidFill', None, E('a:srgbClr', {'val': color})), E('a:effectLst')))
    cSld.insert(0, bg)


def clear_shapes(slide):
    tree = _sp_tree(slide)
    for el in list(tree):
        if el.tag in ('{%s}nvGrpSpPr' % P, '{%s}grpSpPr' % P):
            continue
        tree.remove(el)


# ----------------------------------------------------------------------------- theme blocks
def label(slide, text, dark=False):
    return textbox(slide, 0.62, 0.39, 8.75, 0.23, [para([(text, 11, SKY if dark else BLUE, True)])], name='Section label')


def footer(slide, text, dark=False, sz=10):
    return textbox(slide, 0.62, 5.24, 8.36, 0.2, [para([(text, sz, MIST if dark else GREY)])], name='Footer caveat')


def page_number(slide, n, dark=False):
    return textbox(slide, 9.14, 5.24, 0.39, 0.21, [para([(f'{n:02d}', 10, MIST if dark else GREY)], algn='r')], name='Slide number')


def headline(slide, lines, x=0.59, y=1.17, w=8.98, sz=34, dark=False, accent_from=None, h=None, name='Headline'):
    """Multi-line headline, one paragraph per line. accent_from: index of first line drawn in accent colour."""
    color = WHITE if dark else INK
    acc = SKY if dark else BLUE
    paras = []
    for i, ln in enumerate(lines):
        c = acc if (accent_from is not None and i >= accent_from) else color
        paras.append(para([(ln, sz, c, True)]))
    h = h or (len(lines) * sz * 1.2 / 72 + 0.1)
    return textbox(slide, x, y, w, h, paras, name=name)


# ----------------------------------------------------------------------------- notes
def set_notes(slide, text):
    """Drop any existing notes slide, create a fresh one with `text`."""
    part = slide.part
    for rel in list(part.rels.values()):
        if rel.reltype.endswith('/notesSlide'):
            part.drop_rel(rel.rId)
    ns = slide.notes_slide  # creates a new one from the notes master
    tf = ns.notes_text_frame
    tf.text = text
    return ns


# ----------------------------------------------------------------------------- transitions
def set_transition(slide, dur_ms=500):
    sld = slide._element
    for el in list(sld):
        if el.tag == '{%s}transition' % P or el.tag == '{%s}AlternateContent' % MC:
            sld.remove(el)
    ac = etree.Element('{%s}AlternateContent' % MC, nsmap={'mc': MC, 'p14': P14})
    choice = etree.SubElement(ac, '{%s}Choice' % MC)
    choice.set('Requires', 'p14')
    tr = etree.SubElement(choice, '{%s}transition' % P)
    tr.set('spd', 'fast')
    tr.set('{%s}dur' % P14, str(dur_ms))
    etree.SubElement(tr, '{%s}fade' % P)
    fb = etree.SubElement(ac, '{%s}Fallback' % MC)
    tr2 = etree.SubElement(fb, '{%s}transition' % P)
    tr2.set('spd', 'fast')
    etree.SubElement(tr2, '{%s}fade' % P)
    # insert after clrMapOvr (order: cSld, clrMapOvr, transition, timing, extLst)
    clr = sld.find('{%s}clrMapOvr' % P)
    idx = list(sld).index(clr) + 1 if clr is not None else 1
    sld.insert(idx, ac)


# ----------------------------------------------------------------------------- animations
class Timing:
    """Builds a PowerPoint <p:timing> tree.

    groups: list of click groups. Each group is a list of effects (spid, effect, [delay_ms]).
    The first group may be auto-started (start='auto') so that its effects play "With Previous"
    as soon as the slide appears; every other group starts On Click.
    effect in {'fade', 'wipeLeft', 'flyUpFade'}; durations in ms.
    """
    DUR = {'fade': 500, 'wipeLeft': 400, 'flyUpFade': 500}

    def __init__(self):
        self.groups = []
        self.auto_first = False
        self._id = 0

    def nid(self):
        self._id += 1
        return str(self._id)

    def add(self, *effects, auto=False):
        """effects: (spid, effect[, dur_ms][, delay_ms]) tuples -> one click (or auto) group."""
        items = []
        for e in effects:
            spid, eff = e[0], e[1]
            dur = e[2] if len(e) > 2 and e[2] else self.DUR[eff]
            delay = e[3] if len(e) > 3 else 0
            items.append((str(spid), eff, dur, delay))
        if auto:
            self.auto_first = True
        self.groups.append(items)
        return self

    def _tgt(self, spid):
        return E('p:tgtEl', None, E('p:spTgt', {'spid': spid}))

    def _set_visible(self, spid):
        return E('p:set', None,
                 E('p:cBhvr', None,
                   E('p:cTn', {'id': self.nid(), 'dur': '1', 'fill': 'hold'}, E('p:stCondLst', None, E('p:cond', {'delay': '0'}))),
                   self._tgt(spid),
                   E('p:attrNameLst', None, E('p:attrName', None, text='style.visibility'))),
                 E('p:to', None, E('p:strVal', {'val': 'visible'})))

    def _anim_effect(self, spid, filt, dur):
        return E('p:animEffect', {'transition': 'in', 'filter': filt},
                 E('p:cBhvr', None, E('p:cTn', {'id': self.nid(), 'dur': str(dur)}), self._tgt(spid)))

    def _anim_pos(self, spid, attr, v0, v1, dur):
        return E('p:anim', {'calcmode': 'lin', 'valueType': 'num'},
                 E('p:cBhvr', {'additive': 'base'},
                   E('p:cTn', {'id': self.nid(), 'dur': str(dur), 'fill': 'hold'}),
                   self._tgt(spid),
                   E('p:attrNameLst', None, E('p:attrName', None, text=attr))),
                 E('p:tavLst', None,
                   E('p:tav', {'tm': '0'}, E('p:val', None, E('p:strVal', {'val': v0}))),
                   E('p:tav', {'tm': '100000'}, E('p:val', None, E('p:strVal', {'val': v1})))))

    def _effect(self, spid, eff, dur, delay, node_type):
        preset = {'fade': ('10', '0'), 'wipeLeft': ('22', '8'), 'flyUpFade': ('2', '4')}[eff]
        ctn = E('p:cTn', {'id': self.nid(), 'presetID': preset[0], 'presetClass': 'entr', 'presetSubtype': preset[1],
                          'fill': 'hold', 'grpId': '0', 'nodeType': node_type},
                E('p:stCondLst', None, E('p:cond', {'delay': str(delay)})))
        ch = E('p:childTnLst')
        ch.append(self._set_visible(spid))
        if eff == 'fade':
            ch.append(self._anim_effect(spid, 'fade', dur))
        elif eff == 'wipeLeft':
            ch.append(self._anim_effect(spid, 'wipe(left)', dur))
        elif eff == 'flyUpFade':
            ch.append(self._anim_effect(spid, 'fade', dur))
            ch.append(self._anim_pos(spid, 'ppt_x', '#ppt_x', '#ppt_x', dur))
            ch.append(self._anim_pos(spid, 'ppt_y', '1+#ppt_h/2', '#ppt_y', dur))
        ctn.append(ch)
        return E('p:par', None, ctn)

    def build(self):
        self._id = 0
        root_ctn = E('p:cTn', {'id': self.nid(), 'dur': 'indefinite', 'restart': 'never', 'nodeType': 'tmRoot'})
        main_ctn = E('p:cTn', {'id': self.nid(), 'dur': 'indefinite', 'nodeType': 'mainSeq'})
        main_children = E('p:childTnLst')
        main_ctn.append(main_children)
        for gi, items in enumerate(self.groups):
            auto = (gi == 0 and self.auto_first)
            st = E('p:stCondLst', None, E('p:cond', {'delay': 'indefinite'}))
            if auto:
                st.append(E('p:cond', {'evt': 'onBegin', 'delay': '0'}, E('p:tn', {'val': '2'})))
            g_ctn = E('p:cTn', {'id': self.nid(), 'fill': 'hold'}, st)
            inner_ctn = E('p:cTn', {'id': self.nid(), 'fill': 'hold'}, E('p:stCondLst', None, E('p:cond', {'delay': '0'})))
            inner_children = E('p:childTnLst')
            for ii, (spid, eff, dur, delay) in enumerate(items):
                node_type = 'withEffect' if (auto or ii > 0) else 'clickEffect'
                inner_children.append(self._effect(spid, eff, dur, delay, node_type))
            inner_ctn.append(inner_children)
            g_ctn.append(E('p:childTnLst', None, E('p:par', None, inner_ctn)))
            main_children.append(E('p:par', None, g_ctn))
        seq = E('p:seq', {'concurrent': '1', 'nextAc': 'seek'}, main_ctn,
                E('p:prevCondLst', None, E('p:cond', {'evt': 'onPrev', 'delay': '0'}, E('p:tgtEl', None, E('p:sldTgt')))),
                E('p:nextCondLst', None, E('p:cond', {'evt': 'onNext', 'delay': '0'}, E('p:tgtEl', None, E('p:sldTgt')))))
        root_ctn.append(E('p:childTnLst', None, seq))
        timing = E('p:timing', None, E('p:tnLst', None, E('p:par', None, root_ctn)))
        return timing

    def apply(self, slide):
        sld = slide._element
        for el in list(sld):
            if el.tag == '{%s}timing' % P:
                sld.remove(el)
        if not self.groups:
            return
        timing = self.build()
        # bldLst for text shapes (p:sp with txBody) so PowerPoint treats them as whole-shape builds
        bld = E('p:bldLst')
        seen = set()
        sp_ids = {e.get('id') for e in sld.iter('{%s}cNvPr' % P) if e.getparent().getparent().tag == '{%s}sp' % P}
        for items in self.groups:
            for spid, *_ in items:
                if spid in sp_ids and spid not in seen:
                    seen.add(spid)
                    bld.append(E('p:bldP', {'spid': spid, 'grpId': '0'}))
        if len(bld):
            timing.append(bld)
        # insert after transition/AlternateContent, before extLst
        ext = sld.find('{%s}extLst' % P)
        if ext is not None:
            ext.addprevious(timing)
        else:
            sld.append(timing)
